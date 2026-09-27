"""End-to-end parsing with the real Docling / Unstructured stack on small generated files.

Skipped automatically when the heavy dependencies are missing (e.g. a local 3.13 venv); run
them in the Docker test image:  docker build --target test -t doc-processor:test . &&
docker run --rm doc-processor:test
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

pytest.importorskip("docling")
pytest.importorskip("unstructured")
fpdf = pytest.importorskip("fpdf")

from app.chunking import hf_token_counter  # noqa: E402
from app.config import Settings  # noqa: E402
from app.models import DocumentFailedEvent, DocumentParsedEvent, DocumentUploadedEvent, ErrorCode  # noqa: E402
from app.pipeline import Pipeline, build_pipeline  # noqa: E402
from app.probes import probe_pdf  # noqa: E402
from tests.conftest import upload_event  # noqa: E402

pytestmark = pytest.mark.docling

PARAGRAPH = (
    "A database transaction groups several reads and writes so that they succeed or fail "
    "together. Isolation levels trade consistency for concurrency: stronger levels prevent "
    "anomalies such as dirty reads, non-repeatable reads and phantoms. "
)
TABLE_ROWS = [("Level", "Dirty reads", "Phantoms"), ("Read committed", "No", "Yes"), ("Serializable", "No", "No")]
SCANNED_LINES = [
    "Kafka keeps ordering within a partition.",
    "Consumers commit offsets after processing.",
    "At least once delivery allows duplicates.",
]


def _write_text_pdf(path: Path, pages: int = 3) -> None:
    pdf = fpdf.FPDF()
    for page in range(pages):
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 18)
        pdf.cell(0, 12, f"Chapter {page + 1}: Transactions", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=11)
        for _ in range(5):
            pdf.multi_cell(0, 6, PARAGRAPH * 2, new_x="LMARGIN", new_y="NEXT")
        if page == 1:
            pdf.ln(4)
            with pdf.table() as table:
                for row in TABLE_ROWS:
                    cells = table.row()
                    for value in row:
                        cells.cell(value)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))


def _write_scanned(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1700, 1000), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=44)
    for i, line in enumerate(SCANNED_LINES):
        draw.text((80, 100 + i * 90), line, fill="black", font=font)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


@pytest.fixture(scope="module")
def real_settings(tmp_path_factory: pytest.TempPathFactory) -> Settings:
    return Settings(storage_root=tmp_path_factory.mktemp("storage"), _env_file=None)


@pytest.fixture(scope="module")
def real_pipeline(real_settings: Settings) -> Pipeline:
    return build_pipeline(real_settings)


def _event(pack: str, ext: str, **overrides: object) -> DocumentUploadedEvent:
    return upload_event(int(pack), ext, eventId=f"e-{pack}", tier="MAX", **overrides)


def _chunks(settings: Settings, event: DocumentParsedEvent) -> list[dict]:
    return json.loads((settings.storage_root / event.chunks_path).read_text(encoding="utf-8"))


def test_docling_parses_text_pdf(real_pipeline: Pipeline, real_settings: Settings) -> None:
    _write_text_pdf(real_settings.storage_root / "packs/1/source.pdf")
    result = real_pipeline.process(_event("1", "pdf"))

    assert isinstance(result, DocumentParsedEvent), result
    assert result.parser == "docling" and result.page_count == 3 and result.ocr_used is False
    chunks = _chunks(real_settings, result)
    assert result.chunk_count == len(chunks) > 1
    count = hf_token_counter(real_settings.tokenizer_model)
    for i, chunk in enumerate(chunks):
        assert chunk["index"] == i
        assert chunk["text"].strip()
        assert count(chunk["text"]) <= real_settings.chunk_max_tokens
        assert 1 <= chunk["page"] <= chunk["pageEnd"] <= 3
    assert any((c["section"] or "").startswith("Chapter") for c in chunks)
    tables = [c for c in chunks if c["elementType"] == "table"]
    assert tables and "| Serializable" in tables[0]["text"] and "|---" in tables[0]["text"].replace(" ", "")


def test_page_limit_is_checked_before_parsing(real_pipeline: Pipeline, real_settings: Settings) -> None:
    _write_text_pdf(real_settings.storage_root / "packs/2/source.pdf", pages=4)
    result = real_pipeline.process(_event("2", "pdf", maxPages=3))
    assert isinstance(result, DocumentFailedEvent)
    assert result.error_code == ErrorCode.PAGE_LIMIT_EXCEEDED
    assert result.message == "Document has 4 pages; your tier allows 3."


def test_scanned_pdf_needs_ocr_then_docling_ocr_reads_it(real_pipeline: Pipeline, real_settings: Settings) -> None:
    png = real_settings.storage_root / "tmp/scan.png"
    _write_scanned(png)
    from PIL import Image

    pdf_path = real_settings.storage_root / "packs/3/source.pdf"
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    Image.open(png).save(pdf_path)
    assert probe_pdf(pdf_path, 50).has_text_layer is False

    denied = real_pipeline.process(_event("3", "pdf", ocrEnabled=False))
    assert isinstance(denied, DocumentFailedEvent) and denied.error_code == ErrorCode.OCR_REQUIRED

    ocr = real_pipeline.process(_event("3", "pdf", ocrEnabled=True))
    assert isinstance(ocr, DocumentParsedEvent), ocr
    assert ocr.parser == "docling" and ocr.ocr_used is True
    assert "partition" in " ".join(c["text"] for c in _chunks(real_settings, ocr)).lower()


def test_image_ocr_with_unstructured(real_pipeline: Pipeline, real_settings: Settings) -> None:
    _write_scanned(real_settings.storage_root / "packs/4/source.png")
    result = real_pipeline.process(_event("4", "png", ocrEnabled=True))
    assert isinstance(result, DocumentParsedEvent), result
    assert result.parser == "unstructured" and result.ocr_used and result.page_count == 1
    assert "offsets" in " ".join(c["text"] for c in _chunks(real_settings, result)).lower()


def test_docx_with_headings_and_table(real_pipeline: Pipeline, real_settings: Settings) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Databases", level=1)
    document.add_heading("Transactions", level=2)
    for _ in range(4):
        document.add_paragraph(PARAGRAPH * 2)
    table = document.add_table(rows=2, cols=2)
    for r, row in enumerate([("Level", "Phantoms"), ("Serializable", "No")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_paragraph("• first bullet", style="List Bullet")
    path = real_settings.storage_root / "packs/5/source.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))

    result = real_pipeline.process(_event("5", "docx"))
    assert isinstance(result, DocumentParsedEvent), result
    chunks = _chunks(real_settings, result)
    assert any(c["section"] == "Databases > Transactions" for c in chunks), chunks
    tables = [c for c in chunks if c["elementType"] == "table"]
    assert tables and tables[0]["text"].startswith("| Level | Phantoms |")


def test_pptx_slides(real_pipeline: Pipeline, real_settings: Settings) -> None:
    from pptx import Presentation

    deck = Presentation()
    for i in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i + 1}: Replication"
        slide.placeholders[1].text = "Leaders accept writes; followers replicate the log."
    path = real_settings.storage_root / "packs/6/source.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(path))

    result = real_pipeline.process(_event("6", "pptx"))
    assert isinstance(result, DocumentParsedEvent), result
    assert result.page_count == 3
    chunks = _chunks(real_settings, result)
    assert {c["page"] for c in chunks} <= {1, 2, 3}

    limited = real_pipeline.process(_event("6", "pptx", maxPages=2))
    assert isinstance(limited, DocumentFailedEvent) and limited.error_code == ErrorCode.PAGE_LIMIT_EXCEEDED


def test_unstructured_pdf_fallback(real_settings: Settings) -> None:
    from app.parsers.base import DocKind
    from app.parsers.unstructured_parser import UnstructuredParser

    path = real_settings.storage_root / "packs/7/source.pdf"
    _write_text_pdf(path, pages=2)
    output = UnstructuredParser(real_settings).parse(path, DocKind.PDF, ocr=False)
    assert output.chunks and output.page_count == 2
    assert "transaction" in " ".join(c.text for c in output.chunks).lower()


def test_corrupt_pdf_is_parse_error(real_pipeline: Pipeline, real_settings: Settings) -> None:
    path = real_settings.storage_root / "packs/8/source.pdf"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"%PDF-1.4 this is not really a pdf")
    result = real_pipeline.process(_event("8", "pdf"))
    assert isinstance(result, DocumentFailedEvent) and result.error_code == ErrorCode.PARSE_ERROR
